"""Synthetic, offline workbook contract tests. No company documents or model calls."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import load_workbook
from excel_uret import build, validate
from excel_dogrula import expected_values, parameter_cases, preflight
from veri_kontrol import digest


def fixture():
    return {'schema': 'teklif-data/v4', 'project': 'Sentetik Excel doğrulama', 'analysis_date': '2026-09-26',
            'decision_summary': 'Sonuç ve Öneri: Bu sentetik örnek satınalma kararı içermez.',
            'offers': [{'id': 'A', 'supplier_id': 'SA', 'name': 'Firma A', 'currency': 'EUR',
                        'lines': [{'id': 'L1', 'description': '=UNTRUSTED()', 'quantity': '2', 'unit_price': '120',
                                   'price_unit': '1', 'discounts': ['0.10', '0.05'], 'vat_rate': '0.20', 'price_includes_vat': True}]}],
            'costs': {'A': {'base_date': '2026-09-26', 'annual_rates': {'EUR': '0.04'}, 'base_currency': 'TRY',
                            'fx_rates': {'EUR': {'forex_selling': '40', 'unit': '1'}},
                            'events': [{'event_id': 'offer', 'owner': 'A', 'known': True, 'amount': '1000', 'currency': 'EUR', 'payment_date': '2027-09-26'},
                                       {'event_id': 'unknown', 'owner': 'A', 'known': False, 'currency': 'EUR', 'payment_date': '2027-09-26'}]}},
            'facts': [], 'scope_items': [], 'rfi': [{'id': 'R1', 'owner': 'Firma A', 'question': 'Navlun tutarı nedir?', 'fact_ids': []}]}


def scored_fixture():
    data = fixture(); data['costs']['A']['events'].pop()
    other = copy.deepcopy(data['offers'][0]); other.update(id='B', supplier_id='SB', name='Firma B'); data['offers'].append(other)
    data['costs']['B'] = copy.deepcopy(data['costs']['A']); data['costs']['B']['events'][0]['amount'] = '1200'
    data['scoring'] = {'suppliers': [{'supplier_id': 'SA', 'eligible': True}, {'supplier_id': 'SB', 'eligible': True}],
                       'criteria': [{'criterion_id': 'cost', 'weight': '60', 'direction': 'lower', 'values': {'SA': '1000', 'SB': '1200'}},
                                    {'criterion_id': 'quality', 'weight': '40', 'direction': 'direct', 'values': {'SA': '7', 'SB': '9'}}]}
    data['scoring']['criteria'][0]['bindings'] = {s: {'kind': 'cost', 'offer_id': o, 'metric': 'known_present_value_in_base_currency'} for s, o in [('SA', 'A'), ('SB', 'B')]}
    data['scoring']['criteria'][1]['bindings'] = {s: {'kind': 'fact', 'fact_id': s + '/quality'} for s in ('SA', 'SB')}
    data['scoring']['sensitivity'] = [{'id': 'cost_focus', 'label': 'Maliyet ağırlığı yüksek', 'weights': {'cost': '70', 'quality': '30'}},
                                    {'id': 'quality_focus', 'label': 'Kalite ağırlığı yüksek', 'weights': {'cost': '50', 'quality': '50'}}]
    data['facts'] = [{'fact_id': s + '/quality', 'evidence_status': 'verified', 'value': v} for s, v in [('SA', '7'), ('SB', '9')]]
    data['requirements'] = [{'id': 'REQ1'}]
    return data


def decision_fixture(data):
    return {'data_sha256': digest(data), 'summary': 'Sonuç ve Öneri: Sentetik test; kaynaklar eşit kapsamlı olarak yeniden teyit edilmelidir.', 'recommendation': None}


class ExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='v4-excel-test-'); self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / 'analysis.xlsx'

    def test_profiles_and_no_fabricated_caches(self):
        data = fixture(); result = build(data, self.path, 'hizli')
        self.assertEqual(result['sheets'], ['Özet', 'Karar Özeti', 'Fiyat ve Kapsam', 'RFI'])
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        preflight(self.path, data, contract)
        wb = load_workbook(self.path, data_only=True)
        for output in contract['outputs']: self.assertIsNone(wb[output['sheet']][output['cell']].value)
        wb.close()
        with self.assertRaises(ValueError): build(data, self.path)

    def test_compounded_discount_and_inclusive_vat_match_motor(self):
        data = fixture(); result = build(data, self.path)
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        values = expected_values(data, contract)
        net = next(o for o in contract['outputs'] if o['kind'] == 'line' and o['field'] == 'net_excluding_vat')
        self.assertEqual(str(values[net['sheet'], net['cell']]), '171.00')
        unknown = next(o for o in contract['outputs'] if o['kind'] == 'event_pv' and o['event_id'] == 'unknown')
        self.assertIsNone(values[unknown['sheet'], unknown['cell']])
        wb = load_workbook(self.path)
        self.assertTrue(any(c.data_type == 's' and '=UNTRUSTED()' in c.value for ws in wb for row in ws for c in row if isinstance(c.value, str)))
        wb.close()

    def test_zero_price_and_numeric_json_inputs(self):
        data = fixture(); data['offers'][0]['lines'][0]['unit_price'] = 0
        data['costs']['A']['annual_rates']['EUR'] = 0.04
        validate(data); result = build(data, self.path)
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        self.assertIn(0, [v for k, v in expected_values(data, contract).items() if v is not None])

    def test_scored_standard_has_nine_tabs_and_independent_parameter_cases(self):
        data = scored_fixture(); result = build(data, self.path)
        self.assertEqual(len(result['sheets']), 9)
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        preflight(self.path, data, contract)
        kinds = {c['kind'] for c in parameter_cases(data, contract)}
        self.assertTrue({'amount', 'date', 'rate', 'forex_selling', 'unit', 'weight'}.issubset(kinds))
        cases = parameter_cases(data, contract)
        self.assertTrue(all(len(c['changes']) == 2 for c in cases if c['kind'] in ('rate', 'forex_selling', 'unit')))
        self.assertTrue(any(o['kind'] == 'sensitivity_score' for o in contract['outputs']))

    def test_formula_or_input_tampering_fails_preflight(self):
        data = fixture(); result = build(data, self.path)
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        wb = load_workbook(self.path); o = contract['outputs'][0]; wb[o['sheet']][o['cell']] = '=42'; wb.save(self.path); wb.close()
        with self.assertRaisesRegex(ValueError, 'Formüller'): preflight(self.path, data, contract)

    def test_incomplete_cost_cannot_have_scores(self):
        data = scored_fixture(); data['costs']['A']['events'][0]['known'] = False
        with self.assertRaises(ValueError): validate(data)

    def test_external_decision_is_bound_without_changing_data(self):
        data = fixture(); original = digest(data); decision = decision_fixture(data)
        result = build(data, self.path, decision=decision)
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        self.assertEqual(digest(data), original); self.assertEqual(contract['decision_sha256'], digest(decision))
        preflight(self.path, data, contract, decision)
        changed = dict(decision, summary='different')
        with self.assertRaises(ValueError): preflight(self.path, data, contract, changed)

    def test_freshness_checks_exist_without_validity_or_quote_date(self):
        data = fixture(); result = build(data, self.path, 'hizli')
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        values = expected_values(data, contract)
        self.assertIn('Geçerlilik teyidi yok', values.values()); self.assertIn('Teklif tarihi yok', values.values())

    def test_all_unknown_costs_do_not_turn_into_zero_summary(self):
        data = fixture(); data['costs']['A']['events'][0]['known'] = False
        result = build(data, self.path, 'hizli')
        contract = json.loads(Path(result['contract']).read_text(encoding='utf-8'))
        self.assertFalse(any(o['kind'] in ('nominal', 'pv', 'base_currency_total') for o in contract['outputs']))

    def test_sensitivity_requires_two_distinct_balanced_alternatives(self):
        data = scored_fixture(); data['scoring']['sensitivity'][0]['weights']['cost'] = '71'
        with self.assertRaisesRegex(ValueError, '100'): validate(data)
        data = scored_fixture(); data['scoring']['sensitivity'][0]['weights'] = {'cost': '60', 'quality': '40'}
        with self.assertRaisesRegex(ValueError, 'farklı'): validate(data)

    def test_single_criterion_explains_absence_of_weight_alternatives(self):
        data = scored_fixture(); data['scoring']['criteria'] = data['scoring']['criteria'][:1]
        data['scoring']['criteria'][0]['weight'] = '100'; data['scoring']['sensitivity'] = []
        build(data, self.path)
        wb = load_workbook(self.path)
        self.assertTrue(any('Tek kriterde ağırlık alternatifi yoktur' in str(c.value) for row in wb['Puanlama'] for c in row))
        wb.close()


if __name__ == '__main__': unittest.main()

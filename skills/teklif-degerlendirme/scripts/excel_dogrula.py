"""Excel COM recalculation, independent motor reconciliation and parameter restore.

An isolated worker owns a DispatchEx instance. Existing Excel processes are never
closed. On a deadline, only the recorded, newly created process can be terminated.
The original workbook is replaced only after every check succeeds on a copy.
"""
from __future__ import annotations
import argparse
import contextlib
import copy
import datetime as dt
from decimal import Decimal
import gc
import json
import math
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import uuid

from excel_uret import file_hash, normalized, validate, sensitivity_specs
from teklif_motoru import cost_summary, price_line, score_suppliers
from veri_kontrol import digest, bound_scoring


def at(data, path, default=None):
    try:
        for key in path: data = data[key]
        return data
    except (KeyError, IndexError, TypeError):
        return default


def put(data, path, value):
    obj = data
    for key in path[:-1]: obj = obj[key]
    obj[path[-1]] = value


def expected_values(data, contract):
    """Recompute from the actual central data, never trust contract-stored totals."""
    costs = {k: cost_summary(normalized(p)) for k, p in data.get('costs', {}).items()}
    scoring_input = bound_scoring(normalized(data), costs) if data.get('scoring') else None
    scoring = score_suppliers(scoring_input) if scoring_input else None
    scenarios = {}
    if scoring_input:
        for spec in sensitivity_specs(scoring_input):
            variant = copy.deepcopy(scoring_input)
            for criterion in variant['criteria']: criterion['weight'] = spec['weights'][criterion['criterion_id']]
            scenarios[spec['id']] = score_suppliers(variant)
    sr = {r['supplier_id']: r for r in scoring['suppliers']} if scoring else {}
    values = {}
    for o in contract['outputs']:
        kind = o['kind']; key = (o['sheet'], o['cell']); oid = o.get('offer_id')
        if kind.startswith('event_'):
            event = next((r for r in costs[oid]['events'] if r['event_id'] == o['event_id']), None)
            value = None if event is None else event['present_value']
            if value is not None and kind == 'event_base':
                quote = data['costs'][oid]['fx_rates'][event['currency']]
                value = value * Decimal(str(quote['forex_selling'])) / Decimal(str(quote['unit']))
        elif kind in ('nominal', 'pv', 'base_currency_total'):
            curr = o['currency']
            value = costs[oid]['known_nominal_by_currency' if kind == 'nominal' else 'known_present_value_by_currency'][curr]
            if kind == 'base_currency_total':
                quote = data['costs'][oid]['fx_rates'][curr]
                value = value * Decimal(str(quote['forex_selling'])) / Decimal(str(quote['unit']))
        elif kind == 'line':
            line = data['offers'][o['offer_index']]['lines'][o['line_index']]
            value = price_line(normalized(line))[o['field']] if all(line.get(k) is not None for k in ('quantity', 'unit_price', 'vat_rate')) else None
        elif kind in ('quote_age', 'age_status', 'validity'):
            field = 'validity' if kind == 'validity' else 'quote_date'
            fact = next((f for f in data.get('facts', []) if f['fact_id'] == f'{oid}/field/{field}'), {})
            date = dt.date.fromisoformat(fact['value']) if fact.get('evidence_status') == 'verified' else None
            today = dt.date.fromisoformat(data['analysis_date'])
            if kind == 'validity': value = 'Geçerlilik teyidi yok' if date is None else ('Geçerlilik dolmuş' if date < today else 'Süre içinde')
            elif kind == 'quote_age': value = None if date is None else (today - date).days
            else: value = 'Teklif tarihi yok' if date is None else ('Güncel teklif teyidi' if (today-date).days > 60 or (today-date).days < 0 else 'Yaş sınırı içinde')
        elif kind == 'distinguishing_weight': value = scoring['distinguishing_weight']
        elif kind == 'criterion_value': value = Decimal(next(c for c in scoring_input['criteria'] if c['criterion_id'] == o['criterion_id'])['values'][o['supplier_id']])
        elif kind == 'criterion_score': value = sr[o['supplier_id']]['criterion_scores'][o['criterion_id']]
        elif kind == 'score': value = sr[o['supplier_id']]['total_score']
        elif kind == 'rank': value = sr[o['supplier_id']]['rank']
        elif kind.startswith('sensitivity_'):
            scenario = scenarios[o['scenario_id']]
            if kind == 'sensitivity_weight_total': value = Decimal('100')
            elif kind == 'sensitivity_ratio': value = scenario['distinguishing_weight_ratio']
            else:
                supplier = next(s for s in scenario['suppliers'] if s['supplier_id'] == o['supplier_id'])
                value = supplier['total_score'] if kind == 'sensitivity_score' else supplier['rank']
        else: raise ValueError('Bilinmeyen hesap sözleşmesi: ' + kind)
        values[key] = value
    return values


def equal(actual, expected):
    if expected is None: return actual is None or actual == ''
    if isinstance(expected, str): return actual == expected
    if isinstance(actual, bool) or not isinstance(actual, (int, float, Decimal)): return False
    if not math.isfinite(float(actual)): return False
    # Excel uses IEEE754. Keep sub-cent absolute tolerance and tight relative tolerance.
    return abs(Decimal(str(actual)) - Decimal(str(expected))) <= max(Decimal('0.000001'), abs(Decimal(str(expected))) * Decimal('1e-11'))


def layout_check(workbook):
    """Mechanical wrapping/height and Unicode checks; visual review stays separate."""
    failures = []
    for ws in workbook:
        if ws.max_row > 100000 or ws.max_column > 256:
            failures.append(ws.title + ': düzen sınırı aşıldı'); continue
        if not ws.print_area or not ws.freeze_panes:
            failures.append(ws.title + ': yazdırma alanı/sabit başlık yok')
        for row in ws:
            for cell in row:
                if not isinstance(cell.value, str) or cell.data_type == 'f': continue
                if '\ufffd' in cell.value: failures.append(f'{ws.title}!{cell.coordinate}: bozuk Unicode')
                if any(cell.coordinate in r for r in ws.merged_cells.ranges): continue
                width = ws.column_dimensions[cell.column_letter].width or 22
                lines = sum(max(1, math.ceil(len(x)/max(width-2, 8))) for x in cell.value.splitlines())
                height = ws.row_dimensions[cell.row].height or 15
                if lines * 14 + 8 > height + 1:
                    failures.append(f'{ws.title}!{cell.coordinate}: metin yüksekliği yetersiz')
    return failures


def preflight(path, data, contract=None, decision=None):
    from openpyxl import load_workbook
    from openpyxl.utils.datetime import to_excel
    validate(data)
    if contract is None:
        contract = json.loads(Path(path).with_suffix(Path(path).suffix + '.contract.json').read_text(encoding='utf-8'))
    if contract.get('schema') != 'teklif-excel-contract/v4' or contract.get('data_sha256') != digest(data):
        raise ValueError('Excel sözleşmesi merkezi veriyle eşleşmiyor.')
    if decision is not None and (decision.get('data_sha256') != digest(data) or contract.get('decision_sha256') != digest(decision)):
        raise ValueError('Excel hakem özeti sözleşmesi farklı.')
    if not contract.get('outputs') or not contract.get('formulas'):
        raise ValueError('Bağımsız hesapla denetlenecek sonuç/formül yok.')
    wb = load_workbook(path, data_only=False, keep_links=False)
    try:
        if wb.sheetnames != contract['sheets']: raise ValueError('Sekme sözleşmesi farklı.')
        for c in contract.get('decision_cells', []):
            if wb['Karar Özeti'][c['cell']].value != c['value']:
                raise ValueError('Excel karar metni hakem kaydından farklı.')
        planned = {(f['sheet'], f['cell']): f['formula'] for f in contract['formulas']}
        actual = {(ws.title, c.coordinate): c.value for ws in wb for row in ws for c in row if c.data_type == 'f'}
        if planned != actual: raise ValueError('Formüller üretim sözleşmesinden farklı.')
        for i in contract['inputs']:
            value = at(data, i['path'], i.get('default')); observed = wb[i['sheet']][i['cell']].value
            if i['kind'] == 'date':
                wanted = None if value is None else dt.date.fromisoformat(value)
                if isinstance(observed, dt.datetime): observed = observed.date()
                if observed != wanted: raise ValueError('Tarih girdisi merkezi veriden farklı.')
            elif not equal(observed, None if value is None else Decimal(str(value))):
                raise ValueError('Sayı girdisi merkezi veriden farklı: ' + str(i['path']))
        failures = layout_check(wb)
        if failures: raise ValueError('; '.join(failures[:10]))
    finally: wb.close()
    expected_values(data, contract)
    return True


def parameter_cases(data, contract):
    """One meaningful perturbation per available input family; weights move as a pair."""
    baseline = expected_values(data, contract); cases = []; seen = set()
    for i in contract['inputs']:
        kind = i['kind']; old = at(data, i['path'], i.get('default'))
        if kind in seen or kind in ('weight', 'sensitivity_weight') or old is None: continue
        changed = copy.deepcopy(data)
        if kind == 'date':
            delta = -1 if i['path'][-1] == 'base_date' else 1
            new = (dt.date.fromisoformat(old) + dt.timedelta(days=delta)).isoformat()
        else:
            value = Decimal(str(old))
            if kind in ('rate', 'vat_rate'): new = str(value + Decimal('0.01'))
            elif kind == 'discount': new = str(value - Decimal('0.01') if value >= Decimal('0.98') else value + Decimal('0.01'))
            elif kind == 'criterion': new = str(value - Decimal('0.1') if value >= Decimal('9.9') else value + Decimal('0.1'))
            else: new = str(value + max(Decimal('1'), abs(value) * Decimal('0.1')))
        # Shared market assumptions must move together for every offer. The
        # independent bound_scoring gate rejects inconsistent currency assumptions.
        shared = [i]
        if kind == 'rate': shared = [x for x in contract['inputs'] if x['kind'] == kind and x['path'][-1] == i['path'][-1]]
        elif kind in ('forex_selling', 'unit'): shared = [x for x in contract['inputs'] if x['kind'] == kind and x['path'][-2] == i['path'][-2]]
        elif kind == 'date' and i['path'][-1] == 'base_date': shared = [x for x in contract['inputs'] if x['kind'] == 'date' and x['path'][-1] == 'base_date']
        else: shared = [x for x in contract['inputs'] if x['path'] == i['path']]
        changes = []
        for x in shared:
            put(changed, x['path'], new); changes.append({'input': x, 'value': new})
        try: expected = expected_values(changed, contract)
        except ValueError: continue
        if any(not equal(baseline[k], v) for k, v in expected.items()):
            cases.append({'kind': kind, 'changes': changes, 'data': changed}); seen.add(kind)
    weight_groups = [[i for i in contract['inputs'] if i['kind'] == 'weight']]
    for scenario_index in range(len(data.get('scoring', {}).get('sensitivity', []))):
        weight_groups.append([i for i in contract['inputs'] if i['kind'] == 'sensitivity_weight' and i['path'][2] == scenario_index])
    for weights in weight_groups:
        if len(weights) < 2: continue
        donor = next((i for i in weights if Decimal(str(at(data, i['path']))) > 0), None)
        if donor:
            recipient = next(i for i in weights if i is not donor)
            for step in ('1', '0.1', '0.01'):
                delta = min(Decimal(step), Decimal(str(at(data, donor['path']))))
                changed = copy.deepcopy(data); changes = []
                for i, sign in [(donor, -1), (recipient, 1)]:
                    value = str(Decimal(str(at(data, i['path']))) + sign * delta)
                    put(changed, i['path'], value); changes.append({'input': i, 'value': value})
                try: expected = expected_values(changed, contract)
                except ValueError: continue
                if any(not equal(baseline[k], v) for k, v in expected.items()):
                    cases.append({'kind': weights[0]['kind'], 'changes': changes, 'data': changed}); break
    return cases


def _process_identity(pid):
    import win32api, win32process, win32con
    handle = win32api.OpenProcess(win32con.PROCESS_QUERY_INFORMATION, False, pid)
    try: return str(win32process.GetProcessTimes(handle)['CreationTime'])
    finally: handle.Close()


def _stop_owned(state):
    """Never target by process name, never terminate an existing or reused PID."""
    try:
        import win32api, win32con
        if not state.get('owned') or _process_identity(state['pid']) != state['created']: return False
        handle = win32api.OpenProcess(win32con.PROCESS_TERMINATE, False, state['pid'])
        try: win32api.TerminateProcess(handle, 1)
        finally: handle.Close()
        return True
    except Exception: return False


@contextlib.contextmanager
def _workspace_temp(parent):
    """Inherit the output directory ACL; Windows private temp DACLs must not travel."""
    import shutil
    parent = Path(parent).resolve(); stage = parent / ('.excel-verify-' + uuid.uuid4().hex)
    stage.mkdir(mode=0o755)
    try:
        yield stage
    finally:
        if stage.parent.resolve() != parent or not stage.name.startswith('.excel-verify-') or stage.is_symlink() or (hasattr(stage, 'is_junction') and stage.is_junction()):
            raise ValueError('Geçici Excel dizini sınırı doğrulanamadı.')
        shutil.rmtree(stage)


def _worker(request_path):
    request = json.loads(Path(request_path).read_text(encoding='utf-8'))
    result = {'ok': False, 'checks': {'recalculation': False, 'reconciliation': False, 'parameter_restore': False, 'layout': True}}
    app = book = None; owned = False
    try:
        import pythoncom
        import pywintypes
        import win32com.client
        import win32process
        from openpyxl.utils.datetime import to_excel
        pythoncom.CoInitialize()
        before = set(win32process.EnumProcesses())
        app = win32com.client.DispatchEx('Excel.Application')
        _, pid = win32process.GetWindowThreadProcessId(app.Hwnd)
        if pid in before: raise ValueError('Excel yeni ve bağımsız süreç olarak açılamadı.')
        owned = True
        state = {'pid': pid, 'created': _process_identity(pid), 'owned': True}
        Path(request['state']).write_text(json.dumps(state), encoding='utf-8')
        app.Visible = False; app.DisplayAlerts = False; app.AskToUpdateLinks = False
        app.EnableEvents = False; app.AutomationSecurity = 3
        book = app.Workbooks.Open(str(Path(request['workbook']).resolve()), UpdateLinks=0, ReadOnly=False, IgnoreReadOnlyRecommended=True)
        app.Calculation = -4105  # xlCalculationAutomatic; scoped to this private instance.
        book.ForceFullCalculation = False
        contract = request['contract']; data = request['data']

        def calculate():
            app.CalculateFullRebuild()
            for sheet in book.Worksheets:
                sheet.Calculate()
            until = time.monotonic() + 60
            # xlPending (2) can remain application-wide after all target sheets
            # calculate. Reconcile every target output and its perturbations below;
            # only xlCalculating (1) means calculation is actively unfinished.
            while app.CalculationState == 1:
                if time.monotonic() > until: raise TimeoutError(f'Excel hesaplama süresi aşıldı; durum={app.CalculationState}, mod={app.Calculation}.')
                pythoncom.PumpWaitingMessages(); time.sleep(0.05)
            result['calculation_state'] = app.CalculationState

        def reconcile(expected):
            for (sheet, cell), wanted in expected.items():
                value = book.Worksheets(sheet).Range(cell).Value2
                if not equal(value, wanted): raise ValueError(f'Motor/Excel farkı: {sheet}!{cell}: {value!r} / {wanted!r}')

        calculate()
        baseline = expected_values(data, contract); reconcile(baseline)
        for ws in book.Worksheets:
            # SpecialCells uses enum values, avoiding localized Evaluate names.
            for cell_type in (-4123, 2):  # formulas, constants; 16 = errors
                try:
                    errors = ws.UsedRange.SpecialCells(cell_type, 16)
                except pywintypes.com_error as exc:
                    if exc.excepinfo and exc.excepinfo[5] == -2146827284:
                        continue  # Excel 1004: no matching cells in this valid range.
                    raise
                raise ValueError(f'Excel hata hücresi: {ws.Name}!{errors.Address}')
        result['checks']['recalculation'] = True
        result['checks']['reconciliation'] = True
        cases = parameter_cases(data, contract)
        if not cases: raise ValueError('Sonucu değiştiren anlamlı parametre testi bulunamadı.')
        result['parameter_tests'] = []
        for case in cases:
            originals = []
            try:
                for change in case['changes']:
                    i = change['input']; target = book.Worksheets(i['sheet']).Range(i['cell'])
                    originals.append((i, target.Value2))
                    target.Value2 = to_excel(dt.datetime.combine(dt.date.fromisoformat(change['value']), dt.time())) if i['kind'] == 'date' else float(change['value'])
                calculate(); reconcile(expected_values(case['data'], contract))
            finally:
                for i, value in originals: book.Worksheets(i['sheet']).Range(i['cell']).Value2 = value
                calculate()
            reconcile(baseline); result['parameter_tests'].append({'kind': case['kind'], 'restored': True})
        result['checks']['parameter_restore'] = True
        book.Save()
        if request.get('pdf'):
            book.ExportAsFixedFormat(0, str(Path(request['pdf']).resolve()))
        book.Close(SaveChanges=False); book = None
        result['ok'] = True
    except Exception as error:
        result['error'] = f'{type(error).__name__}: {error}'
    finally:
        if book is not None and owned:
            try: book.Close(SaveChanges=False)
            except Exception: pass
        book = None
        if app is not None and owned:
            try: app.Quit()
            except Exception: pass
        app = None; gc.collect()
        try: pythoncom.CoUninitialize()
        except Exception: pass
        Path(request['result']).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return 0 if result['ok'] else 2


def verify(workbook, data, *, profile=None, pdf=None, timeout=180, contract_path=None, decision=None):
    """Return PASS only after COM calculation; unavailable Excel is UNVERIFIED."""
    import shutil
    workbook = Path(workbook).resolve()
    cp = Path(contract_path) if contract_path else workbook.with_suffix(workbook.suffix + '.contract.json')
    receipt_path = workbook.with_suffix(workbook.suffix + '.verification.json')
    original_sha = file_hash(workbook)
    receipt = {'schema': 'teklif-excel-verification/v4', 'status': 'UNVERIFIED', 'workbook_sha256': original_sha,
               'data_sha256': digest(data), 'contract_sha256': file_hash(cp), 'decision_sha256': digest(decision) if decision is not None else None,
               'checks': {'recalculation': False, 'reconciliation': False, 'parameter_restore': False, 'layout': False},
               'checked_at': dt.datetime.now(dt.timezone.utc).isoformat(), 'visual_review': 'separate_required_for_high_assurance'}
    try:
        contract = json.loads(cp.read_text(encoding='utf-8'))
        if contract.get('decision_sha256') != receipt['decision_sha256']:
            raise ValueError('Excel doğrulamasına aynı hakem özeti verilmelidir.')
        if profile is not None and profile != contract['profile']: raise ValueError('Profil sözleşmesi farklı.')
        if receipt_path.exists():
            previous = json.loads(receipt_path.read_text(encoding='utf-8'))
            keys = ('workbook_sha256', 'data_sha256', 'contract_sha256', 'decision_sha256')
            if previous.get('status') == 'PASS' and all(previous.get(k) == receipt[k] for k in keys):
                if not contract['pdf_required'] or (previous.get('pdf_path') and Path(previous['pdf_path']).is_file() and file_hash(previous['pdf_path']) == previous.get('pdf_sha256')):
                    return {**previous, 'reused': True}
        if original_sha != contract['workbook_sha256']:
            raise ValueError('Üretilen kitap değişmiş; yeni sürüm üretin.')
        preflight(workbook, data, contract, decision); receipt['checks']['layout'] = True
        pdf_target = Path(pdf).resolve() if pdf else (workbook.with_suffix('.pdf') if contract['pdf_required'] else None)
        if pdf_target and pdf_target.exists(): raise ValueError('Var olan PDF üzerine yazılmaz.')
        with _workspace_temp(workbook.parent) as folder:
            work = Path(folder); copybook = work / 'analysis.xlsx'; shutil.copy2(workbook, copybook)
            request = {'workbook': str(copybook), 'data': data, 'contract': contract, 'state': str(work / 'excel-owner.json'),
                       'result': str(work / 'result.json'), 'pdf': str(work / 'report.pdf') if pdf_target else None}
            reqpath = work / 'request.json'; reqpath.write_text(json.dumps(request, ensure_ascii=False), encoding='utf-8')
            try:
                run = subprocess.run([sys.executable, '-B', str(Path(__file__).resolve()), '--worker', str(reqpath)],
                                     capture_output=True, text=True, encoding='utf-8', timeout=timeout, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
            except subprocess.TimeoutExpired:
                state = json.loads(Path(request['state']).read_text(encoding='utf-8')) if Path(request['state']).exists() else {}
                _stop_owned(state); raise TimeoutError('Excel doğrulama zaman sınırına ulaştı; yalnız bu doğrulamaya ait süreç kapatıldı.')
            if not Path(request['result']).is_file(): raise ValueError('Excel işçisi makbuz üretmedi: ' + run.stderr[-500:])
            result = json.loads(Path(request['result']).read_text(encoding='utf-8'))
            receipt['checks'] = result['checks']; receipt['parameter_tests'] = result.get('parameter_tests', [])
            receipt['calculation_state'] = result.get('calculation_state')
            state_path = Path(request['state'])
            if state_path.exists():
                state = json.loads(state_path.read_text(encoding='utf-8'))
                alive = False
                for _ in range(20):
                    try: alive = _process_identity(state['pid']) == state['created']
                    except Exception: alive = False
                    if not alive: break
                    time.sleep(0.1)
                receipt['owned_process_forced_close'] = _stop_owned(state) if alive else False
                receipt['excel_process_closed'] = not alive or receipt['owned_process_forced_close']
                if not receipt['excel_process_closed']: raise ValueError('Doğrulamaya ait Excel süreci kapanmadı.')
            if not result.get('ok'): raise ValueError(result.get('error', 'Excel kontrolü tamamlanmadı.'))
            # Read the saved cache again: a successful in-memory COM evaluation is insufficient.
            from openpyxl import load_workbook
            cached = load_workbook(copybook, data_only=True, keep_links=False)
            try:
                for (sheet, cell), wanted in expected_values(data, contract).items():
                    if not equal(cached[sheet][cell].value, wanted): raise ValueError('Kaydedilen Excel önbelleği motorla farklı.')
            finally: cached.close()
            if pdf_target:
                from cikti_denetimi import validate_pdf
                validate_pdf(request['pdf'])
            if file_hash(workbook) != original_sha: raise ValueError('Doğrulama sırasında kaynak kitap değişti.')
            os.replace(copybook, workbook)
            if pdf_target:
                pdf_target.parent.mkdir(parents=True, exist_ok=True); os.replace(request['pdf'], pdf_target)
                receipt['pdf_sha256'] = file_hash(pdf_target); receipt['pdf_path'] = str(pdf_target)
            receipt['workbook_sha256'] = file_hash(workbook); receipt['status'] = 'PASS'
    except Exception as error:
        receipt['error'] = f'{type(error).__name__}: {error}'
        receipt['status'] = 'UNVERIFIED'
    receipt_path.write_text(json.dumps(receipt, ensure_ascii=False, indent=2), encoding='utf-8')
    return receipt


if __name__ == '__main__':
    if hasattr(sys.stdout, 'reconfigure'): sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser(description=__doc__); parser.add_argument('--worker', type=Path)
    parser.add_argument('--workbook', type=Path); parser.add_argument('--data', type=Path); parser.add_argument('--pdf', type=Path); parser.add_argument('--profile'); parser.add_argument('--decision', type=Path)
    args = parser.parse_args()
    if args.worker: raise SystemExit(_worker(args.worker))
    if not args.workbook or not args.data: parser.error('--workbook ve --data gerekli.')
    decision = json.loads(args.decision.read_text(encoding='utf-8')) if args.decision else None
    result = verify(args.workbook, json.loads(args.data.read_text(encoding='utf-8')), profile=args.profile, pdf=args.pdf, decision=decision)
    print(json.dumps(result, ensure_ascii=False)); raise SystemExit(0 if result['status'] == 'PASS' else 2)

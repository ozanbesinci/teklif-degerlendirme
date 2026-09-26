"""Read-only environment and document inspection. Never installs software."""
from __future__ import annotations
import argparse
import importlib.metadata as md
import json
from pathlib import Path
import sys
import zipfile
import xml.etree.ElementTree as ET
from email import policy
from email.parser import BytesParser
from kayit_temeli import sha

VERSIONS = {"openpyxl": "3.1.5", "pypdf": "6.13.2", "pdfplumber": "0.11.10", "pypdfium2": "5.10.1", "pywin32": "312"}


def environment():
    packages, errors = {}, []
    for name, expected in VERSIONS.items():
        try:
            actual = md.version(name)
        except md.PackageNotFoundError:
            actual = None
        packages[name] = {"expected": expected, "actual": actual}
        if actual != expected:
            errors.append(f"{name}: beklenen {expected}, kurulu {actual}")
    if sys.version_info[:2] != (3, 14):
        errors.append("program-guncelle ile desteklenen Python 3.14 gerekli.")
    return {"status": "PASS" if not errors else "UNVERIFIED", "python": sys.version,
            "executable": sys.executable, "packages": packages, "issues": errors,
            "remedy": "Eksik/uyumsuz ortam için program-guncelle kullanın; burada kurulum yapılmaz." if errors else None}


def extract(path, output):
    path, output = Path(path).resolve(), Path(output).resolve()
    if output.exists() or output == path or path.is_relative_to(output):
        raise ValueError("Yeni ve kaynaktan ayrı çıkarım dizini gerekli.")
    output.mkdir(parents=True)
    ext = path.suffix.lower()
    report = {"source": str(path), "source_sha256":sha(path), "status": "READ", "parts": []}
    if ext == ".pdf":
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for i, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                part = {"page": i, "text_file": f"page-{i}.txt", "needs_visual": not text.strip()}
                (output / part["text_file"]).write_text(text, encoding="utf-8")
                if part["needs_visual"]:
                    import pypdfium2 as pdfium
                    document = pdfium.PdfDocument(path)
                    p = document[i-1]
                    bitmap = p.render(scale=1.5)
                    image = bitmap.to_pil()
                    image.save(output / f"page-{i}.png")
                    image.close(); bitmap.close(); p.close(); document.close()
                    part["image_file"] = f"page-{i}.png"
                report["parts"].append(part)
    elif ext in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        try:
            for i, sheet in enumerate(wb, 1):
                part = {"sheet": sheet.title, "state": sheet.sheet_state, "text_file": f"sheet-{i}.json",
                        "hidden_rows": [n for n, dim in sheet.row_dimensions.items() if dim.hidden],
                        "hidden_columns": [n for n, dim in sheet.column_dimensions.items() if dim.hidden]}
                rows = [[{"cell": c.coordinate, "value": c.value} for c in row if c.value is not None] for row in sheet]
                (output / part["text_file"]).write_text(json.dumps(rows, default=str, ensure_ascii=False), encoding="utf-8")
                report["parts"].append(part)
        finally:
            wb.close()
    elif ext == ".docx":
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith("word/") and name.endswith(".xml") and any(k in name for k in ("document", "footnote", "endnote", "header", "footer")):
                    root = ET.fromstring(archive.read(name))
                    text = "\n".join("".join(p.itertext()) for p in root.iter(ns+"p"))
                    filename = name.replace("/", "_") + ".txt"
                    (output/filename).write_text(text, encoding="utf-8")
                    report["parts"].append({"part": name, "text_file": filename,
                                            "tracked_changes": any(True for _ in root.iter(ns+"del")) or any(True for _ in root.iter(ns+"ins"))})
        report["layout_warning"] = "OOXML metni sayfa düzeni değildir; izlenen değişiklikler ve kritik tablolar özgün dosyada doğrulanır."
    elif ext == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        attachments = []
        for index, part in enumerate(message.walk()):
            if part.get_content_disposition() == "attachment":
                attachments.append({"name": part.get_filename(), "content_type": part.get_content_type(), "needs_read": True})
        text = "\n".join(str(message.get(k, "")) for k in ("From", "To", "Date", "Subject"))
        body = message.get_body(preferencelist=("plain", "html"))
        if body:
            text += "\n" + str(body.get_content())
        (output/"message.txt").write_text(text, encoding="utf-8")
        report["parts"] = [{"text_file": "message.txt", "attachments": attachments}]
        if attachments:
            report["status"] = "PARTIAL"
    elif ext in {".txt", ".csv"}:
        text = path.read_text(encoding="utf-8-sig")
        (output/"text.txt").write_text(text, encoding="utf-8")
        report["parts"] = [{"text_file": "text.txt"}]
    else:
        report.update(status="UNVERIFIED", reason="Bu biçim otomatik çözümlenmedi; özgün dosya uygun araçla okunmalı. Sessizce dışlanmaz.")
    (output/"extraction.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def render_pages(path, output, pages):
    """Render supplied 1-based original PDF pages, including text-bearing pages."""
    import pypdfium2 as pdfium
    path,output=Path(path).resolve(),Path(output).resolve()
    if output.exists() or path.is_relative_to(output): raise ValueError('Yeni görüntü dizini gerekli.')
    if not pages or len(set(pages))!=len(pages) or any(type(p) is not int or p<1 for p in pages):
        raise ValueError('Benzersiz 1 tabanlı sayfa listesi gerekli.')
    output.mkdir(parents=True); rows=[]
    with pdfium.PdfDocument(path) as doc:
        if max(pages)>len(doc): raise ValueError('Sayfa belge dışında.')
        for number in pages:
            page=doc[number-1]; bitmap=page.render(scale=1.5); image=bitmap.to_pil()
            target=output/f'page-{number}.png'; image.save(target)
            image.close(); bitmap.close(); page.close()
            rows.append({'page':number,'image_path':str(target),'image_sha256':sha(target)})
    report={'source':str(path),'source_sha256':sha(path),'images':rows}
    (output/'images.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source"); parser.add_argument("--output")
    parser.add_argument('--pages',help='Virgülle ayrılmış özgün PDF sayfalarını görüntüle (1 tabanlı).')
    args = parser.parse_args()
    if args.pages and (not args.source or not args.output): parser.error('--pages için source ve output gerekli')
    result = render_pages(args.source,args.output,[int(p) for p in args.pages.split(',')]) if args.pages else extract(args.source, args.output) if args.source and args.output else environment()
    print(json.dumps(result, ensure_ascii=False))
